from __future__ import annotations

import argparse
import json
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.neighbors import NearestNeighbors

from hre_translate.data.normalize import normalize_hre, normalize_vietnamese
from hre_translate.data.preprocess import read_sheet
from hre_translate.data.split import similarity_document
from hre_translate.utils.config import load_config, resolve


def _json_value(value: Any) -> Any:
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    return value


def _length_stats(series: pd.Series) -> dict[str, float | int]:
    lengths = series.dropna().astype(str).str.len()
    if lengths.empty:
        return {"min": 0, "mean": 0.0, "median": 0.0, "p95": 0.0, "max": 0}
    return {
        "min": int(lengths.min()),
        "mean": round(float(lengths.mean()), 2),
        "median": float(lengths.median()),
        "p95": float(lengths.quantile(0.95)),
        "max": int(lengths.max()),
    }


def _characters(series: pd.Series) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    counter = Counter("".join(series.dropna().astype(str)))
    top = [{"character": char, "count": count} for char, count in counter.most_common(40)]
    suspicious = []
    for char, count in counter.items():
        category = unicodedata.category(char)
        if char == "\ufffd" or category in {"Cc", "Cf", "Cs", "Co", "Cn"}:
            suspicious.append(
                {
                    "character": char,
                    "codepoint": f"U+{ord(char):04X}",
                    "name": unicodedata.name(char, "UNKNOWN"),
                    "category": category,
                    "count": count,
                }
            )
    return top, sorted(suspicious, key=lambda item: item["count"], reverse=True)


def _near_duplicates(frame: pd.DataFrame, threshold: float, limit: int = 100) -> dict[str, Any]:
    clean = frame.dropna(subset=["hre", "vi"]).reset_index(drop=True)
    if len(clean) < 2:
        return {"threshold": threshold, "pair_count": 0, "examples": []}
    docs = [
        similarity_document(record.hre, record.vi)
        for record in clean[["hre", "vi"]].itertuples(index=False)
    ]
    matrix = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit_transform(docs)
    model = NearestNeighbors(metric="cosine", algorithm="brute", radius=1.0 - threshold).fit(matrix)
    distances, neighbors = model.radius_neighbors(matrix, return_distance=True)
    examples: list[dict[str, Any]] = []
    pair_count = 0
    for left, (row_distances, row_neighbors) in enumerate(zip(distances, neighbors, strict=True)):
        for distance, right in zip(row_distances, row_neighbors, strict=True):
            right = int(right)
            if right <= left:
                continue
            pair_count += 1
            if len(examples) < limit:
                examples.append(
                    {
                        "left": {"hre": clean.at[left, "hre"], "vi": clean.at[left, "vi"]},
                        "right": {"hre": clean.at[right, "hre"], "vi": clean.at[right, "vi"]},
                        "similarity": round(1.0 - float(distance), 4),
                    }
                )
    examples.sort(key=lambda item: item["similarity"], reverse=True)
    return {"threshold": threshold, "pair_count": pair_count, "examples": examples}


def _sheet_report(
    raw: pd.DataFrame,
    spec: dict[str, Any],
    threshold: float,
) -> dict[str, Any]:
    hre_raw, vi_raw = raw[spec["hre_column"]], raw[spec["vi_column"]]
    normalized = pd.DataFrame(
        {"hre": hre_raw.map(normalize_hre), "vi": vi_raw.map(normalize_vietnamese)}
    )
    paired = normalized.dropna(subset=["hre", "vi"])
    top_hre, suspicious_hre = _characters(hre_raw)
    top_vi, suspicious_vi = _characters(vi_raw)
    word_pairs = int(
        (paired["hre"].str.split().str.len().eq(1) & paired["vi"].str.split().str.len().eq(1)).sum()
    )
    return {
        "columns": [str(column) for column in raw.columns],
        "row_count": len(raw),
        "null_counts": {str(key): int(value) for key, value in raw.isna().sum().items()},
        "usable_pair_count": len(paired),
        "empty_or_incomplete_pair_count": len(raw) - len(paired),
        "exact_duplicate_pair_count": int(paired.duplicated(["hre", "vi"]).sum()),
        "source_length_characters": _length_stats(normalized["hre"]),
        "target_length_characters": _length_stats(normalized["vi"]),
        "source_vocabulary_size": len(
            {token for text in normalized["hre"].dropna() for token in text.split()}
        ),
        "target_vocabulary_size": len(
            {token for text in normalized["vi"].dropna() for token in text.split()}
        ),
        "top_source_characters": top_hre,
        "top_target_characters": top_vi,
        "suspicious_source_characters": suspicious_hre,
        "suspicious_target_characters": suspicious_vi,
        "word_pair_count": word_pairs,
        "phrase_pair_count": len(paired) - word_pairs,
        "possible_near_duplicates": _near_duplicates(paired, threshold),
    }


def _markdown(report: dict[str, Any]) -> str:
    lines = ["# Data inspection report", "", "## Files", ""]
    for item in report["files"]:
        lines.append(f"- `{item['filename']}`: {item['type']}, {item['size_bytes']:,} bytes")
    lines.extend(["", "## Workbook", ""])
    for name, sheet in report["workbook"]["sheets"].items():
        lines.extend(
            [
                f"### {name}",
                "",
                f"- Rows read: {sheet['row_count']:,}",
                f"- Usable pairs: {sheet['usable_pair_count']:,}",
                f"- Exact duplicates: {sheet['exact_duplicate_pair_count']:,}",
                f"- Word pairs: {sheet['word_pair_count']:,}",
                f"- Phrase pairs: {sheet['phrase_pair_count']:,}",
                "- Possible near-duplicate pairs: "
                f"{sheet['possible_near_duplicates']['pair_count']:,}",
                "",
            ]
        )
    archive = report["audio_archive"]
    lines.extend(
        [
            "## Audio archive",
            "",
            f"- MP3 members: {archive['mp3_count']:,}",
            f"- Total uncompressed bytes: {archive['uncompressed_bytes']:,}",
            "",
        ]
    )
    return "\n".join(lines)


def run_inspection(config_path: str | Path | None = None) -> dict[str, Any]:
    config, root = load_config(config_path)
    workbook = resolve(root, config["raw"]["workbook"])
    archive = resolve(root, config["raw"]["audio_archive"])
    threshold = float(config["processing"]["near_duplicate_threshold"])
    workbook_meta = load_workbook(workbook, read_only=True, data_only=True)
    sheets: dict[str, Any] = {}
    for section in ("lexicon", "sentences"):
        spec = config["raw"][section]
        sheets[spec["sheet"]] = _sheet_report(read_sheet(workbook, spec), spec, threshold)
        sheets[spec["sheet"]]["excel_max_row"] = workbook_meta[spec["sheet"]].max_row
        sheets[spec["sheet"]]["excel_max_column"] = workbook_meta[spec["sheet"]].max_column

    with zipfile.ZipFile(archive) as handle:
        members = [item for item in handle.infolist() if not item.is_dir()]
        mp3 = [item for item in members if item.filename.lower().endswith(".mp3")]
    report = {
        "files": [
            {
                "filename": workbook.name,
                "type": "Excel workbook",
                "size_bytes": workbook.stat().st_size,
            },
            {
                "filename": archive.name,
                "type": "ZIP audio archive",
                "size_bytes": archive.stat().st_size,
            },
        ],
        "workbook": {
            "filename": workbook.name,
            "sheet_names": workbook_meta.sheetnames,
            "sheets": sheets,
        },
        "audio_archive": {
            "filename": archive.name,
            "member_count": len(members),
            "mp3_count": len(mp3),
            "uncompressed_bytes": sum(item.file_size for item in members),
            "min_numeric_stem": min(int(Path(item.filename).stem) for item in mp3),
            "max_numeric_stem": max(int(Path(item.filename).stem) for item in mp3),
        },
        "dataset_counts": {
            "word_pairs": sheets[config["raw"]["lexicon"]["sheet"]]["word_pair_count"],
            "phrase_pairs": sheets[config["raw"]["lexicon"]["sheet"]]["phrase_pair_count"],
            "sentence_pairs": sheets[config["raw"]["sentences"]["sheet"]]["usable_pair_count"],
        },
    }
    report = json.loads(json.dumps(report, ensure_ascii=False, default=_json_value))
    json_path = resolve(root, config["outputs"]["report_json"])
    markdown_path = resolve(root, config["outputs"]["report_markdown"])
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    markdown_path.write_text(_markdown(report), encoding="utf-8")
    print(json.dumps(report["dataset_counts"], ensure_ascii=False, indent=2))
    print(f"Reports written to {json_path.relative_to(root)} and {markdown_path.relative_to(root)}")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect raw H're-Vietnamese data.")
    parser.add_argument("--config", type=Path)
    args = parser.parse_args()
    run_inspection(args.config)


if __name__ == "__main__":
    main()
